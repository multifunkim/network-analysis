#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Step 7 — Atom QC and clean k-hubness reconstruction

This step is intentionally separate from Steps 1–6 because it
requires human review.

Workflow
--------
1. Step 6 generates atom files, original k-hubness, and an
   atom_qc_<subject>.xlsx workbook.
2. User labels every saved atom as:
       clean
       noisy
3. Step 7 reads completed QC workbooks.
4. Only atoms labeled "clean" are included in the new k-hubness.
5. The original k-hubness map is NEVER modified.

Supported atom formats
----------------------
NIfTI:
    .nii
    .nii.gz

GIFTI:
    .func.gii
"""

import os
import re
import argparse
import logging
from pathlib import Path

import numpy as np
import nibabel as nib

from openpyxl import load_workbook

from spatial_io import save_gifti_map


# ============================================================
# LOGGING
# ============================================================

def setup_logger():

    logging.basicConfig(
        level=logging.INFO,
        format="[step7] %(asctime)s - %(levelname)s - %(message)s"
    )


# ============================================================
# FIND QC WORKBOOKS
# ============================================================

def find_qc_workbooks(qc_path):
    """
    qc_path may be:
        - one atom_qc_*.xlsx file
        - one KMAP/subject directory
        - a root SPARK results directory
    """

    qc_path = Path(qc_path).resolve()

    if not qc_path.exists():
        raise FileNotFoundError(
            f"QC path does not exist: {qc_path}"
        )

    if qc_path.is_file():

        if (
            qc_path.suffix.lower() == ".xlsx"
            and qc_path.name.startswith("atom_qc_")
        ):
            return [qc_path]

        raise ValueError(
            "Single-file QC input must be an "
            "atom_qc_*.xlsx workbook."
        )

    workbooks = sorted(
        qc_path.rglob(
            "atom_qc_*.xlsx"
        )
    )

    return workbooks


# ============================================================
# READ / VALIDATE WORKBOOK
# ============================================================

def read_qc_workbook(workbook_path):

    wb = load_workbook(
        workbook_path,
        data_only=True
    )

    if "Atoms" not in wb.sheetnames:
        raise ValueError(
            f"{workbook_path}: missing 'Atoms' sheet."
        )

    ws = wb["Atoms"]

    header = [
        cell.value
        for cell in ws[1]
    ]

    required = [
        "atom_id",
        "atom_file",
        "network",
        "label",
        "notes",
        "n_spatial_elements"
    ]

    missing = [
        col
        for col in required
        if col not in header
    ]

    if missing:
        raise ValueError(
            f"{workbook_path}: missing columns {missing}"
        )

    col = {
        name: header.index(name)
        for name in required
    }

    rows = []

    for row in ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        atom_file = row[
            col["atom_file"]
        ]

        if atom_file is None:
            continue

        label = row[
            col["label"]
        ]

        if label is None:
            label = ""
        else:
            label = str(label).strip().lower()

        rows.append({
            "atom_id": row[col["atom_id"]],
            "atom_file": str(atom_file).strip(),
            "network": row[col["network"]],
            "label": label,
            "notes": row[col["notes"]],
        })

    return rows


# ============================================================
# QC STATUS
# ============================================================

def validate_qc(rows):

    if not rows:
        return False, "workbook contains no atoms"

    labels = [
        row["label"]
        for row in rows
    ]

    # Entire workbook untouched
    if all(
        label == ""
        for label in labels
    ):
        return False, "no QC labels entered"

    invalid = sorted(
        {
            label
            for label in labels
            if label not in (
                "",
                "clean",
                "noisy"
            )
        }
    )

    if invalid:
        return False, (
            "invalid QC label(s): "
            + ", ".join(invalid)
        )

    incomplete = sum(
        label == ""
        for label in labels
    )

    if incomplete > 0:
        return False, (
            f"QC incomplete: {incomplete}/{len(rows)} "
            "atoms are unlabeled"
        )

    clean_count = sum(
        row["label"] == "clean"
        for row in rows
    )

    noisy_count = sum(
        row["label"] == "noisy"
        for row in rows
    )

    if clean_count == 0:
        return False, (
            "QC complete but no atoms are labeled clean"
        )

    return True, (
        f"{clean_count} clean / "
        f"{noisy_count} noisy"
    )


# ============================================================
# SUBJECT LABEL
# ============================================================

def subject_from_workbook(path):

    name = Path(path).name

    m = re.match(
        r"atom_qc_(.+)\.xlsx$",
        name
    )

    if not m:
        raise ValueError(
            f"Cannot determine subject from {name}"
        )

    return m.group(1)


# ============================================================
# NIFTI CLEAN K-HUBNESS
# ============================================================

def build_nifti_kmap(
    clean_atom_paths,
    output_path
):

    reference = nib.load(
        str(clean_atom_paths[0])
    )

    ref_shape = reference.shape

    hub = np.zeros(
        ref_shape,
        dtype=np.int32
    )

    for atom_path in clean_atom_paths:

        img = nib.load(
            str(atom_path)
        )

        if img.shape != ref_shape:
            raise ValueError(
                f"Shape mismatch: {atom_path}"
            )

        data = np.asarray(
            img.dataobj
        )

        hub += (
            np.isfinite(data)
            & (data != 0)
        ).astype(
            np.int32
        )

    output_img = nib.Nifti1Image(
        hub.astype(
            np.int16
        ),
        reference.affine,
        reference.header.copy()
    )

    output_img.header.set_data_dtype(
        np.int16
    )

    output_img.header["scl_slope"] = 1.0
    output_img.header["scl_inter"] = 0.0

    qform, qcode = reference.get_qform(
        coded=True
    )

    sform, scode = reference.get_sform(
        coded=True
    )

    if qform is not None:
        output_img.set_qform(
            qform,
            code=int(qcode)
        )

    if sform is not None:
        output_img.set_sform(
            sform,
            code=int(scode)
        )

    nib.save(
        output_img,
        str(output_path)
    )


# ============================================================
# GIFTI CLEAN K-HUBNESS
# ============================================================

def load_gifti_vector(path):

    img = nib.load(
        str(path)
    )

    if len(img.darrays) != 1:
        raise ValueError(
            f"Expected one GIFTI DataArray: {path}"
        )

    return np.asarray(
        img.darrays[0].data
    ).squeeze()


def build_gifti_kmap(
    clean_atom_paths,
    output_path
):

    first = clean_atom_paths[0]

    first_values = load_gifti_vector(
        first
    )

    hub = np.zeros(
        first_values.shape,
        dtype=np.float32
    )

    for atom_path in clean_atom_paths:

        data = load_gifti_vector(
            atom_path
        )

        if data.shape != hub.shape:
            raise ValueError(
                f"Vertex-count mismatch: {atom_path}"
            )

        hub += (
            np.isfinite(data)
            & (data != 0)
        ).astype(
            np.float32
        )

    save_gifti_map(
        values=hub,
        output_path=str(output_path),
        template_path=str(first)
    )


# ============================================================
# PROCESS ONE WORKBOOK
# ============================================================

def process_workbook(
    workbook_path,
    overwrite=False
):

    workbook_path = Path(
        workbook_path
    )

    subject = subject_from_workbook(
        workbook_path
    )

    rows = read_qc_workbook(
        workbook_path
    )

    valid, status = validate_qc(
        rows
    )

    if not valid:

        logging.info(
            "SKIP %s: %s",
            subject,
            status
        )

        return "skipped"

    logging.info(
        "%s: %s",
        subject,
        status
    )

    kmap_dir = workbook_path.parent

    # --------------------------------------------------------
    # Validate atom files
    # --------------------------------------------------------

    clean_atom_paths = []

    for row in rows:

        atom_path = (
            kmap_dir
            / row["atom_file"]
        )

        if not atom_path.is_file():
            raise FileNotFoundError(
                f"{subject}: atom file missing: "
                f"{atom_path}"
            )

        if row["label"] == "clean":
            clean_atom_paths.append(
                atom_path
            )

    if not clean_atom_paths:

        logging.info(
            "SKIP %s: no clean atoms",
            subject
        )

        return "skipped"

    # --------------------------------------------------------
    # Detect format
    # --------------------------------------------------------

    first_atom = clean_atom_paths[0]

    if first_atom.name.endswith(
        ".func.gii"
    ):

        output_path = (
            kmap_dir
            / f"k_hubness_clean_{subject}.func.gii"
        )

        format_name = "gifti"

    elif first_atom.name.endswith(
        ".nii.gz"
    ):

        output_path = (
            kmap_dir
            / f"k_hubness_clean_{subject}.nii.gz"
        )

        format_name = "nifti"

    elif first_atom.name.endswith(
        ".nii"
    ):

        output_path = (
            kmap_dir
            / f"k_hubness_clean_{subject}.nii.gz"
        )

        format_name = "nifti"

    else:

        raise ValueError(
            f"Unsupported atom format: {first_atom}"
        )

    # --------------------------------------------------------
    # Do not overwrite by default
    # --------------------------------------------------------

    if output_path.exists() and not overwrite:

        logging.info(
            "SKIP %s: clean k-hubness already exists: %s",
            subject,
            output_path
        )

        return "skipped"

    # --------------------------------------------------------
    # Reconstruct k-hubness
    # --------------------------------------------------------

    if format_name == "nifti":

        build_nifti_kmap(
            clean_atom_paths,
            output_path
        )

    else:

        build_gifti_kmap(
            clean_atom_paths,
            output_path
        )

    logging.info(
        "DONE %s: %d clean atoms → %s",
        subject,
        len(clean_atom_paths),
        output_path
    )

    return "processed"


# ============================================================
# MAIN
# ============================================================

def main():

    ap = argparse.ArgumentParser(
        description=(
            "Step 7 – Recalculate SPARK k-hubness "
            "from manually QC-labeled atoms"
        )
    )

    ap.add_argument(
        "--qc_path",
        required=True,
        help=(
            "QC workbook, KMAP/subject directory, "
            "or root SPARK results directory"
        )
    )

    ap.add_argument(
        "--overwrite",
        action="store_true",
        help=(
            "Allow replacement of an existing "
            "k_hubness_clean output"
        )
    )

    args = ap.parse_args()

    setup_logger()

    workbooks = find_qc_workbooks(
        args.qc_path
    )

    if not workbooks:

        logging.warning(
            "No atom_qc_*.xlsx workbooks found under %s",
            args.qc_path
        )

        return

    logging.info(
        "Found %d QC workbook(s)",
        len(workbooks)
    )

    processed = 0
    skipped = 0
    failed = 0

    for workbook in workbooks:

        try:

            result = process_workbook(
                workbook_path=workbook,
                overwrite=args.overwrite
            )

            if result == "processed":
                processed += 1
            else:
                skipped += 1

        except Exception as exc:

            failed += 1

            logging.exception(
                "FAILED %s: %s",
                workbook,
                exc
            )

    logging.info(
        "QC summary: processed=%d skipped=%d failed=%d",
        processed,
        skipped,
        failed
    )


if __name__ == "__main__":
    main()
